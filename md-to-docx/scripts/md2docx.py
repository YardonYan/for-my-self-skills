# -*- coding: utf-8 -*-
"""MD -> DOCX 无损转换器（通用，命令行参数化）

用法:
    python md2docx.py <src.md> [dst.docx]
    （dst 省略时 = 同目录同名 .docx）

元素映射规则（无损 = 内容不丢 + 视觉不降级）:
  标题  #~######        -> Heading 1~6（微软雅黑加粗，黑色，保留大纲结构）
  表格  | a | b |       -> Table Grid（表头加粗 + 浅蓝底纹，单元格 9pt，行不跨页）
  代码块 ```            -> 逐行等宽（Consolas + 宋体，浅灰底纹，行距 1.0，缩进）
  引用  >               -> 左竖线 + 浅底纹 + 缩进段
  无序列表 -            -> 项目符号 + 分级缩进（嵌套保留层级）
  有序列表 1.           -> 保留原文编号 + 悬挂缩进
  水平线 ---            -> 细灰线
  行内 `code` / **bold** / *italic* -> 对应 run 格式（数字间星号不做斜体）

依赖: python-docx（pip install python-docx）
验证: 转换后运行 md2docx_verify.py <src.md> <dst.docx> 做零缺失校验
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CN = "微软雅黑"
EN = "Calibri"
CODE_EN = "Consolas"
CODE_CN = "宋体"

HEADING_SIZE = {1: 16, 2: 14, 3: 12, 4: 11, 5: 10.5, 6: 10.5}

# `code` / **bold** / *italic*（斜体要求两侧非字母数字，防 111*84*30cm 误判 *84* 为斜体）
INLINE_RE = re.compile(
    r'(`[^`]+`|\*\*[^*]+\*\*|(?<![A-Za-z0-9])\*[^*]+\*(?![A-Za-z0-9]))'
)


def set_run(run, size=10.5, bold=False, italic=False, cn=CN, en=EN):
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def add_plain(p, text, size=10.5, bold=False, cn=CN, en=EN):
    if not text:
        return
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, cn=cn, en=en)


def add_inline(p, text, size=10.5, base_bold=False):
    """行内格式解析：`code` **bold** *italic*"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            add_plain(p, text[pos:m.start()], size=size, bold=base_bold)
        tok = m.group(0)
        if tok.startswith('`') and tok.endswith('`'):
            r = p.add_run(tok[1:-1])
            set_run(r, size=size - 1, bold=base_bold, cn=CODE_CN, en=CODE_EN)
        elif tok.startswith('**'):
            r = p.add_run(tok[2:-2])
            set_run(r, size=size, bold=True, cn=CN, en=EN)
        else:
            r = p.add_run(tok[1:-1])
            set_run(r, size=size, bold=base_bold, italic=True, cn=CN, en=EN)
        pos = m.end()
    if pos < len(text):
        add_plain(p, text[pos:], size=size, bold=base_bold)


def add_shading(pPr, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def add_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    add_shading(tcPr, fill)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style='Heading %d' % level)
    add_inline(p, text, size=HEADING_SIZE.get(level, 11), base_bold=True)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    return p


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def add_table(doc, header, rows):
    n_cols = max(len(header), max((len(r) for r in rows), default=0))
    n_cols = max(n_cols, 1)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row in table.rows:  # 行不跨页
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    for j in range(n_cols):  # 表头
        cell = table.cell(0, j)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, header[j] if j < len(header) else '', size=9, base_bold=True)
        add_cell_shading(cell, 'D9E2F3')
    for i, row in enumerate(rows, start=1):  # 数据行
        for j in range(n_cols):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, row[j] if j < len(row) else '', size=9)
    return table


def add_code_block(doc, lines):
    n = len(lines)
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(4 if idx == 0 else 0)
        pf.space_after = Pt(4 if idx == n - 1 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.2)
        pf.right_indent = Cm(0.2)
        r = p.add_run(line if line else '')
        set_run(r, size=9, cn=CODE_CN, en=CODE_EN)
        add_shading(p._p.get_or_add_pPr(), 'F5F5F5')


def add_quote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.75)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    add_inline(p, text, size=10.5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), '2E74B5')
    pBdr.append(left)
    pPr.append(pBdr)
    add_shading(pPr, 'F7F9FC')


def add_hr(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'BFBFBF')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(doc, text, indent=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.63 * (indent + 1))
    pf.space_after = Pt(3)
    r = p.add_run('\u2022 ' if indent == 0 else '\u25E6 ')
    set_run(r, size=10.5)
    add_inline(p, text, size=10.5)


def add_ordered(doc, num, text, indent=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.63 * (indent + 1) + 0.5)
    pf.first_line_indent = Cm(-0.5)
    pf.space_after = Pt(3)
    r = p.add_run('%s. ' % num)
    set_run(r, size=10.5)
    add_inline(p, text, size=10.5)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    add_inline(p, text, size=10.5)


SEP_RE = re.compile(r'^\|[\s\-:|]+\|$')


def convert(src, dst):
    with open(src, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)   # A4（默认模板是 Letter，必须改）
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    normal = doc.styles['Normal']
    normal.font.name = EN
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), CN)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith('```'):  # 代码块
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code)
            continue
        if s.startswith('|'):  # 表格
            rows = [split_row(s)]
            i += 1
            if i < n and SEP_RE.match(lines[i].strip()):
                i += 1
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i].strip()))
                i += 1
            add_table(doc, rows[0], rows[1:])
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', line)  # 标题
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue
        if re.match(r'^-{3,}$', s):  # 水平线
            add_hr(doc)
            i += 1
            continue
        if s.startswith('>'):  # 引用
            add_quote(doc, s.lstrip('>').strip())
            i += 1
            continue
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)  # 无序列表
        if m:
            add_bullet(doc, m.group(2).strip(), len(m.group(1)) // 2)
            i += 1
            continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)  # 有序列表
        if m:
            add_ordered(doc, m.group(2), m.group(3).strip(), len(m.group(1)) // 2)
            i += 1
            continue
        add_body(doc, s)  # 普通段落
        i += 1

    doc.save(dst)
    print('OK saved -> %s' % dst)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.docx'
    convert(src, dst)


if __name__ == '__main__':
    main()
